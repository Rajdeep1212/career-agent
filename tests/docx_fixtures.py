"""Synthetic .docx CVs built at test time (no binary fixtures, no real CV content)."""
import io
import re
import struct
import zipfile

import docx
from docx.oxml import parse_xml

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def _bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def plain_cv() -> bytes:
    document = docx.Document()
    for line in ('Asha Verma', 'Education', 'B.Tech in Computer Science, 2025', 'Skills', 'Python, SQL, FastAPI'):
        document.add_paragraph(line)
    return _bytes(document)


def table_cv() -> bytes:
    """Common Word CV template: section headings and details laid out in a two-column table,
    with a vertically merged cell (whose text must not be repeated)."""
    document = docx.Document()
    document.add_paragraph('Rohan Das')
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = 'Education'
    table.cell(0, 1).text = 'B.Tech in Information Technology, 2025'
    table.cell(1, 0).text = 'Skills'
    table.cell(1, 1).text = 'Python, PyTorch, LangChain'
    merged = table.cell(1, 0).merge(table.cell(2, 0))
    merged.text = 'Skills'
    table.cell(2, 1).text = 'SQL, Docker'
    return _bytes(document)


def interleaved_cv() -> bytes:
    """Paragraphs and tables alternate; the text must come out in document order."""
    document = docx.Document()
    document.add_paragraph('Kiran Rao')
    document.add_paragraph('Education')
    first = document.add_table(rows=1, cols=2)
    first.cell(0, 0).text = 'B.Tech in Mechanical Engineering, 2025'
    first.cell(0, 1).text = 'CGPA 8.1'
    document.add_paragraph('Skills')
    second = document.add_table(rows=1, cols=1)
    second.cell(0, 0).text = 'Python, MATLAB'
    document.add_paragraph('Projects')
    return _bytes(document)


def nested_table_cv() -> bytes:
    document = docx.Document()
    document.add_paragraph('Divya Iyer')
    outer = document.add_table(rows=1, cols=2)
    outer.cell(0, 0).text = 'Skills'
    inner = outer.cell(0, 1).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = 'Python, Pandas'
    inner.cell(0, 1).text = 'Tableau, Excel'
    document.add_paragraph('Projects')
    return _bytes(document)


def hyperlink_cv() -> bytes:
    document = docx.Document()
    document.add_paragraph('Sahil Khan')
    paragraph = document.add_paragraph('Portfolio: ')
    paragraph._p.append(parse_xml(
        f'<w:hyperlink xmlns:w="{W}" xmlns:r="{R}" r:id="rId99">'
        '<w:r><w:t>github.com/sahil-example</w:t></w:r></w:hyperlink>'))
    return _bytes(document)


def content_control_cv() -> bytes:
    """Many Word templates wrap body paragraphs in content controls (w:sdt)."""
    document = docx.Document()
    document.add_paragraph('Neha Gupta')
    document.add_paragraph('Projects')
    body = document.element.body
    body.insert(1, parse_xml(
        f'<w:sdt xmlns:w="{W}"><w:sdtContent>'
        '<w:p><w:r><w:t>Summary</w:t></w:r></w:p>'
        '</w:sdtContent></w:sdt>'))
    return _bytes(document)


def header_footer_cv() -> bytes:
    """Name and contact in the header, a note in the footer, over two sections. The second
    section's header is a separate part with the same text (as Word writes for unlinked copies),
    and an empty first-page header must be skipped."""
    document = docx.Document()
    first = document.sections[0]
    first.header.paragraphs[0].text = 'Priya Sen'
    first.header.add_paragraph('priya@example.com')
    first.footer.paragraphs[0].text = 'References available on request'
    first.different_first_page_header_footer = True
    first.first_page_header.paragraphs[0].text = ''
    document.add_paragraph('Education')
    document.add_paragraph('B.Sc in Statistics, 2025')
    second = document.add_section()
    second.header.is_linked_to_previous = False
    second.header.paragraphs[0].text = 'Priya Sen'
    second.header.add_paragraph('priya@example.com')
    document.add_paragraph('Skills')
    document.add_paragraph('R, Python, SQL')
    return _bytes(document)


_TEXT_BOX = (
    '<w:p><w:r><mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
    '<mc:Choice xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" Requires="wps">'
    '<w:drawing><wps:wsp><wps:txbx><w:txbxContent>'
    '<w:p><w:r><w:t>Skills</w:t></w:r></w:p><w:p><w:r><w:t>Python, TensorFlow, RAG</w:t></w:r></w:p>'
    '</w:txbxContent></wps:txbx></wps:wsp></w:drawing></mc:Choice>'
    '<mc:Fallback><w:pict xmlns:v="urn:schemas-microsoft-com:vml"><v:shape><v:textbox><w:txbxContent>'
    '<w:p><w:r><w:t>Skills</w:t></w:r></w:p><w:p><w:r><w:t>Python, TensorFlow, RAG</w:t></w:r></w:p>'
    '</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>'
    '</mc:AlternateContent></w:r></w:p>'
)


def _replace_part(content: bytes, name: str, change) -> bytes:
    source, target = zipfile.ZipFile(io.BytesIO(content)), io.BytesIO()
    with zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as out:
        for item in source.infolist():
            data = source.read(item.filename)
            out.writestr(item, change(data) if item.filename == name else data)
    return target.getvalue()


def text_box_cv() -> bytes:
    """Skills in a sidebar text box, stored twice by Word (modern shape plus VML fallback)."""
    document = docx.Document()
    document.add_paragraph('Meera Nair')
    document.add_paragraph('Education')
    document.add_paragraph('B.Tech in Electronics, 2025')
    content = _bytes(document)
    return _replace_part(content, 'word/document.xml',
                         lambda xml: re.sub(rb'(<w:sectPr)', _TEXT_BOX.encode() + rb'\1', xml, count=1))


_PAYLOAD = 'aaaaaaaaaa'


def entity_expansion_cv() -> bytes:
    """A "billion laughs" document: entity definitions that would expand to gigabytes."""
    lol = (f'<?xml version="1.0"?><!DOCTYPE w:document [<!ENTITY a "{_PAYLOAD}">'
           + ''.join(f'<!ENTITY {chr(98 + i)} "{("&" + chr(97 + i) + ";") * 10}">' for i in range(8))
           + f']><w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>&i;</w:t></w:r></w:p></w:body></w:document>')
    return _replace_part(plain_cv(), 'word/document.xml', lambda _xml: lol.encode())


def external_entity_cv(target_uri: str) -> bytes:
    """An XXE document: an external entity pointing at a local file."""
    xxe = (f'<?xml version="1.0"?><!DOCTYPE w:document [<!ENTITY secret SYSTEM "{target_uri}">]>'
           f'<w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>&secret;</w:t></w:r></w:p></w:body></w:document>')
    return _replace_part(plain_cv(), 'word/document.xml', lambda _xml: xxe.encode())


def zip_bomb_cv() -> bytes:
    """A valid .docx plus one highly compressible entry whose declared size is far past the cap."""
    target = io.BytesIO(plain_cv())
    with zipfile.ZipFile(target, 'a', zipfile.ZIP_DEFLATED) as out:
        out.writestr('word/media/padding.bin', b'\0' * (60 * 1024 * 1024))
    return target.getvalue()


def high_ratio_cv(size: int) -> bytes:
    """Under the declared-size cap, but one entry compresses far beyond any real Word part."""
    target = io.BytesIO(plain_cv())
    with zipfile.ZipFile(target, 'a', zipfile.ZIP_DEFLATED) as out:
        out.writestr('word/media/padding.bin', b'\0' * size)
    return target.getvalue()


def many_entries_cv(count: int) -> bytes:
    target = io.BytesIO(plain_cv())
    with zipfile.ZipFile(target, 'a', zipfile.ZIP_STORED) as out:
        for index in range(count):
            out.writestr(f'extra/item{index}.xml', b'<a/>')
    return target.getvalue()


def lying_size_cv(real_bytes: int, declared_bytes: int = 1000) -> bytes:
    """word/document.xml really inflates to `real_bytes`, but both zip headers claim `declared_bytes`."""
    padding = '<w:p><w:r><w:t>' + 'x' * real_bytes + '</w:t></w:r></w:p>'
    content = _replace_part(plain_cv(), 'word/document.xml',
                            lambda xml: xml.replace(b'<w:sectPr', padding.encode() + b'<w:sectPr', 1))
    archive = zipfile.ZipFile(io.BytesIO(content))
    info = archive.getinfo('word/document.xml')
    raw = bytearray(content)
    # Local file header: uncompressed size at offset 22.
    struct.pack_into('<I', raw, info.header_offset + 22, declared_bytes)
    # Central directory entry: uncompressed size at offset 24.
    central = raw.find(b'PK\x01\x02')
    name = info.filename.encode()
    while central != -1:
        name_length = struct.unpack_from('<H', raw, central + 28)[0]
        if raw[central + 46:central + 46 + name_length] == name:
            struct.pack_into('<I', raw, central + 24, declared_bytes)
            break
        central = raw.find(b'PK\x01\x02', central + 4)
    return bytes(raw)


def bad_zip() -> bytes:
    return b'PK\x03\x04 this is not really a zip archive' + b'\0' * 64


def zip_without_word_document() -> bytes:
    target = io.BytesIO()
    with zipfile.ZipFile(target, 'w') as out:
        out.writestr('[Content_Types].xml', '<Types/>')
        out.writestr('notes.txt', 'not a Word document')
    return target.getvalue()
